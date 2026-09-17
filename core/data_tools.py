import csv
import io
from datetime import date, datetime
from pathlib import Path

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from accounts.models import Role, User
from content.models import News
from core.models import AnalyticsEvent
from core.validators import validate_secure_upload
from feedback.models import Complaint
from issuances.models import Issuance
from offices.models import District, Office
from schools.models import School


class ImportValidationError(Exception):
    def __init__(self, errors):
        self.errors = errors
        super().__init__("Import validation failed")


MAX_IMPORT_ROWS = 5000


def read_rows(upload):
    extension = Path(upload.name).suffix.lower()
    upload.seek(0)
    if extension == ".csv":
        text = upload.read().decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)
    elif extension == ".xlsx":
        workbook = load_workbook(upload, read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(values, [])]
        rows = [dict(zip(headers, row)) for row in values]
    else:
        raise ImportValidationError(["Only CSV and XLSX files are supported."])
    normalized = [
        {str(key).strip().lower(): value for key, value in row.items()}
        for row in rows
        if any(value not in (None, "") for value in row.values())
    ]
    if len(normalized) > MAX_IMPORT_ROWS:
        raise ImportValidationError([f"A single import may contain at most {MAX_IMPORT_ROWS} data rows."])
    return normalized


def _required_headers(rows, required):
    headers = set(rows[0]) if rows else set()
    missing = [name for name in required if name not in headers]
    if missing:
        raise ImportValidationError([f"Missing required column: {name}" for name in missing])


def _integer(value, field, row_number, errors, required=False):
    if value in (None, ""):
        if required:
            errors.append(f"Row {row_number}: {field} is required.")
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        errors.append(f"Row {row_number}: {field} must be a whole number.")
        return None


def _text(value):
    """Normalize spreadsheet values without turning numeric School IDs into 123.0."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _date(value, field, row_number, errors):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value).strip(), "%Y-%m-%d").date()
    except (TypeError, ValueError):
        errors.append(f"Row {row_number}: {field} must use YYYY-MM-DD format.")
        return None


@transaction.atomic
def import_schools(upload):
    rows = read_rows(upload)
    _required_headers(rows, {"school_id", "district", "municipality"})
    errors, prepared, seen = [], [], set()
    districts = {item.name.lower(): item for item in District.objects.filter(active=True)}
    existing = {item.school_id.lower(): item for item in School.objects.all()}
    for number, row in enumerate(rows, 2):
        school_id = _text(row.get("school_id"))
        current = existing.get(school_id.lower())
        name = _text(row.get("name")) or (current.name if current else "")
        district_name = _text(row.get("district"))
        classification = _text(row.get("classification")).upper() or (current.classification if current else "")
        municipality = _text(row.get("municipality"))
        if not school_id or not name:
            errors.append(f"Row {number}: school_id and name are required for new records.")
        if school_id.lower() in seen:
            errors.append(f"Row {number}: duplicate school ID {school_id}.")
        district = districts.get(district_name.lower())
        if not district:
            errors.append(f"Row {number}: unknown or inactive district {district_name}.")
        if classification not in {"PUBLIC", "PRIVATE"}:
            errors.append(f"Row {number}: classification must be PUBLIC or PRIVATE.")
        if not municipality:
            errors.append(f"Row {number}: municipality is required.")
        seen.add(school_id.lower())
        def keep_or_row(field):
            value = _text(row.get(field))
            return value if value else getattr(current, field, "")

        student_population = _integer(row.get("student_population"), "student_population", number, errors)
        teacher_population = _integer(row.get("teacher_population"), "teacher_population", number, errors)
        candidate = School(
            school_id=school_id, name=name, district=district,
            classification=classification, municipality=municipality,
            school_type=keep_or_row("school_type"), level=keep_or_row("level"),
            school_head=keep_or_row("school_head"), address=keep_or_row("address"),
            barangay=keep_or_row("barangay"), contact_number=keep_or_row("contact_number"),
            email=keep_or_row("email"), latitude=row.get("latitude") or getattr(current, "latitude", None),
            longitude=row.get("longitude") or getattr(current, "longitude", None),
            status=_text(row.get("status")).upper() or getattr(current, "status", "ACTIVE"),
            student_population=student_population if student_population is not None else getattr(current, "student_population", None),
            teacher_population=teacher_population if teacher_population is not None else getattr(current, "teacher_population", None),
        )
        prepared.append((current, candidate))
    if errors:
        raise ImportValidationError(errors)
    for _, item in prepared:
        try:
            item.full_clean(validate_unique=False)
        except ValidationError as exc:
            errors.extend(f"{item.school_id}: {message}" for messages in exc.message_dict.values() for message in messages)
    if errors:
        raise ImportValidationError(errors)
    created = updated = 0
    editable_fields = ("name", "district", "classification", "municipality", "school_type", "level", "school_head", "address", "barangay", "contact_number", "email", "latitude", "longitude", "status", "student_population", "teacher_population")
    for current, candidate in prepared:
        if current is None:
            candidate.save()
            created += 1
            continue
        changed = []
        for field in editable_fields:
            value = getattr(candidate, field)
            if getattr(current, field) != value:
                setattr(current, field, value)
                changed.append(field)
        if changed:
            current.save(update_fields=[*changed, "updated_at"])
            updated += 1
    return {"created": created, "updated": updated}


@transaction.atomic
def import_issuances(upload, pdf_files, user):
    rows = read_rows(upload)
    _required_headers(rows, {"reference_number", "title", "category", "year", "date_issued", "office_code", "pdf_filename"})
    errors, prepared, seen = [], [], set()
    offices = {item.code.lower(): item for item in Office.objects.filter(active=True)}
    existing = {value.lower() for value in Issuance.objects.values_list("reference_number", flat=True)}
    pdf_map = {}
    for pdf in pdf_files:
        name = Path(pdf.name).name.lower()
        if name in pdf_map:
            errors.append(f"Duplicate uploaded PDF filename: {name}.")
        try:
            validate_secure_upload(pdf)
            signature = pdf.read(5)
            pdf.seek(0)
            if Path(pdf.name).suffix.lower() != ".pdf" or signature != b"%PDF-":
                errors.append(f"{pdf.name} is not a valid PDF file.")
        except ValidationError as exc:
            errors.extend(f"{pdf.name}: {message}" for message in exc.messages)
        pdf_map[name] = pdf
    valid_categories = {value for value, _ in Issuance.Category.choices}
    for number, row in enumerate(rows, 2):
        reference = str(row.get("reference_number") or "").strip()
        category = str(row.get("category") or "").strip().upper()
        office_code = str(row.get("office_code") or "").strip().lower()
        pdf_name = Path(str(row.get("pdf_filename") or "")).name.lower()
        issued = _date(row.get("date_issued"), "date_issued", number, errors)
        year = _integer(row.get("year"), "year", number, errors, required=True)
        if not reference or not str(row.get("title") or "").strip():
            errors.append(f"Row {number}: reference_number and title are required.")
        if reference.lower() in existing or reference.lower() in seen:
            errors.append(f"Row {number}: duplicate reference number {reference}.")
        if category not in valid_categories:
            errors.append(f"Row {number}: invalid issuance category {category}.")
        office = offices.get(office_code)
        if not office:
            errors.append(f"Row {number}: unknown or inactive office code {office_code}.")
        pdf = pdf_map.get(pdf_name)
        if not pdf:
            errors.append(f"Row {number}: PDF {pdf_name or '(blank)'} was not attached.")
        seen.add(reference.lower())
        prepared.append((row, reference, category, year, issued, office, pdf))
    if errors:
        raise ImportValidationError(errors)
    for row, reference, category, year, issued, office, pdf in prepared:
        pdf.seek(0)
        Issuance.objects.create(
            reference_number=reference, title=str(row["title"]).strip(), category=category,
            year=year, date_issued=issued, office=office, pdf=pdf,
            description="", keywords=str(row.get("keywords") or "").strip(),
            status="DRAFT", created_by=user,
        )
    return len(prepared)


def write_csv(headers, rows):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerows([_safe_row(row) for row in rows])
    return output.getvalue().encode("utf-8-sig")


def write_xlsx(headers, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(_safe_row(row))
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _safe_row(row):
    safe = []
    for value in row:
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            value = "'" + value
        safe.append(value)
    return safe


def write_pdf(title, headers, rows):
    lines = [title, " | ".join(headers), "-" * 105]
    lines.extend(" | ".join(str(value or "") for value in row)[:150] for row in rows)
    pages = [lines[index:index + 48] for index in range(0, len(lines), 48)] or [[title]]
    font_ref = 3 + len(pages) * 2
    objects = [None, b"<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{3 + index * 2} 0 R" for index in range(len(pages)))
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    for index, page_lines in enumerate(pages):
        page_ref, content_ref = 3 + index * 2, 4 + index * 2
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Resources << /Font << /F1 {font_ref} 0 R >> >> /Contents {content_ref} 0 R >>".encode())
        commands = ["BT /F1 8 Tf 36 806 Td"]
        for line in page_lines:
            escaped = str(line).replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands.append(f"({escaped}) Tj 0 -15 Td")
        commands.append("ET")
        stream = "\n".join(commands).encode("latin-1", "replace")
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    document = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects[1:], 1):
        offsets.append(len(document))
        document.extend(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = len(document)
    document.extend(f"xref\n0 {len(objects)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        document.extend(f"{offset:010d} 00000 n \n".encode())
    document.extend(f"trailer << /Size {len(objects)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return bytes(document)


def report_data(report, date_from=None, date_to=None):
    if report == "issuances":
        queryset = Issuance.objects.select_related("office").order_by("-date_issued")
        if date_from: queryset = queryset.filter(date_issued__gte=date_from)
        if date_to: queryset = queryset.filter(date_issued__lte=date_to)
        return ["Reference", "Title", "Category", "Office", "Issued", "Status"], [(x.reference_number, x.title, x.get_category_display(), x.office or "", x.date_issued, x.status) for x in queryset]
    if report == "schools":
        queryset = School.objects.select_related("district").order_by("name")
        return ["School ID", "Name", "District", "Municipality", "Level", "Status"], [(x.school_id, x.name, x.district, x.municipality, x.level, x.status) for x in queryset]
    if report == "news":
        queryset = News.objects.select_related("office", "author").order_by("-created_at")
        if date_from: queryset = queryset.filter(created_at__date__gte=date_from)
        if date_to: queryset = queryset.filter(created_at__date__lte=date_to)
        return ["Title", "Office", "Author", "Published", "Created"], [(x.title, x.office or "", x.author or "", x.is_published, x.created_at.date()) for x in queryset]
    if report == "complaints":
        queryset = Complaint.objects.order_by("-created_at")
        if date_from: queryset = queryset.filter(created_at__date__gte=date_from)
        if date_to: queryset = queryset.filter(created_at__date__lte=date_to)
        return ["Reference", "Subject", "Status", "Created"], [(x.reference_number, x.subject, x.status, x.created_at.date()) for x in queryset]
    if report == "staff":
        queryset = User.objects.filter(is_active=True, role__in=[Role.SUPER_ADMIN, Role.ADMIN, Role.STAFF]).select_related("staff_profile__office")
        return ["Username", "Name", "Role", "Office", "Last login"], [(x.username, x.get_full_name(), x.get_role_display(), getattr(getattr(x, "staff_profile", None), "office", "") or "", x.last_login or "Never") for x in queryset]
    queryset = AnalyticsEvent.objects.order_by("-created_at")
    if date_from: queryset = queryset.filter(created_at__date__gte=date_from)
    if date_to: queryset = queryset.filter(created_at__date__lte=date_to)
    return ["Event", "Object", "Query", "Date"], [(x.get_event_type_display(), f"{x.object_type} {x.object_id}".strip(), x.query, timezone.localtime(x.created_at).strftime("%Y-%m-%d %H:%M")) for x in queryset]
